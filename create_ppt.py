from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
import os

def create_presentation():
    prs = Presentation()
    
    # Title Slide Layout
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "AquaSync AI"
    subtitle.text = "Precision Climate Intelligence for Water-Optimized Agriculture\n\nIndividual Participant | Environmental Innovation Challenge"
    
    # Define a helper for content slides
    def add_slide(title_text, bullet_points):
        layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(layout)
        title = slide.shapes.title
        title.text = title_text
        
        body_shape = slide.shapes.placeholders[1]
        tf = body_shape.text_frame
        
        for i, point in enumerate(bullet_points):
            if i == 0:
                p = tf.paragraphs[0]
                p.text = point
            else:
                p = tf.add_paragraph()
                p.text = point
                p.level = 0
                
    # Slide 2
    add_slide("The Problem: Agriculture is Running Dry", 
              ["70% of global freshwater withdrawals are used for agriculture.",
               "Much of it is wasted through over-irrigation.",
               "Farmers irrigate based on routine and guesswork without localized data.",
               "Results in depleted aquifers and high pumping emissions."])

    # Slide 3
    add_slide("Why Now? Climate Volatility", 
              ["Rising temperatures increase crop water demand unpredictably.",
               "Erratic rainfall makes traditional irrigation schedules obsolete.",
               "Groundwater levels are dropping at alarming rates."])

    # Slide 4
    add_slide("Our Solution: Introducing AquaSync AI", 
              ["A predictive AI platform.",
               "Tells farmers exactly how much water crops need today.",
               "Based on real-time weather forecasts and soil data."])

    # Slide 5
    add_slide("How It Works: Architecture", 
              ["1. Data Sources: Hyper-local weather & soil characteristics.",
               "2. ML Engine: Random Forest predicts exact mm of water.",
               "3. Recommendation: Actionable advice generated.",
               "4. Impact: Saves water, reduces pumping energy."])

    # Slide 6
    add_slide("The AI Engine: Predictive Intelligence", 
              ["Model: Random Forest Regressor",
               "Trained on agronomic climate data.",
               "Predicts exact daily irrigation requirement.",
               "Classifies early warning for crop water stress risk."])

    # Slide 7
    add_slide("Live Prototype", 
              ["1. The Farm Analysis Input screen.",
               "2. The AI Recommendation output (e.g., 'Postpone irrigation - Rain Expected').",
               "(Video to be inserted here)"])

    # Slide 8
    add_slide("Innovation & USPs", 
              ["Predictive, not reactive.",
               "Explainable AI: Tells the farmer WHY the recommendation was made.",
               "Actionable Impact: Automatically calculates water and CO2 saved."])

    # Slide 9
    add_slide("Environmental Impact: Measurable Sustainability", 
              ["Water Saved: Up to 1.5M Liters / Hectare / Year",
               "Emissions: Prevents 1.2 Tons of CO2 / Hectare",
               "Soil Health: Prevents nutrient leaching."])

    # Slide 10
    add_slide("Scalability Roadmap", 
              ["Phase 1: Working ML Prototype (Completed)",
               "Phase 2: Pilot with local farm cooperatives.",
               "Phase 3: API Integration with IoT soil moisture sensors.",
               "Phase 4: Satellite imaging integration."])

    # Slide 11
    add_slide("Securing the Future of Food & Water", 
              ["From predicting climate risk to preventing resource loss.",
               "We cannot afford to waste tomorrow's water using yesterday's methods.",
               "Let's make every drop count."])

    output_path = os.path.join(os.path.dirname(__file__), 'presentation', 'AquaSync_AI_Pitch.pptx')
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == '__main__':
    create_presentation()
